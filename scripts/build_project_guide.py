"""Build the verified SIH26139 project-and-pitch guide as a Word document."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "artifacts" / "docs" / "Q-TRACE_SIH26139_Project_and_Judge_Guide.docx"
CV_PATH = ROOT / "artifacts" / "evaluation" / "five_fold_cross_validation.json"
CLASSICAL_PATH = ROOT / "artifacts" / "evaluation" / "classical_three_seed_metrics.json"

INDIGO = "4A3F8C"
TEAL = "0E7C7B"
INK = "12181C"
MUTED = "647078"
LIGHT = "F2F4F7"
INDIGO_TINT = "F2F0F8"
TEAL_TINT = "EBF5F4"
AMBER_TINT = "FBF5EA"
AMBER = "7A5A00"
WHITE = "FFFFFF"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run(run, *, size: float | None = None, color: str = INK, bold: bool | None = None, italic: bool | None = None, font: str = "Calibri") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.color.rgb = rgb(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table) -> None:
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_DXA:
        raise ValueError(f"table widths must total {CONTENT_DXA} DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag, value in (("w:tblW", CONTENT_DXA), ("w:tblInd", TABLE_INDENT_DXA)):
        element = tbl_pr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_pr.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(table)


def style_table(table, *, font_size: float = 8.5, header_fill: str = LIGHT) -> None:
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            for cell in row.cells:
                set_cell_fill(cell, header_fill)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run(run, size=font_size, bold=(row_index == 0))


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], *, font_size: float = 8.5, header_fill: str = LIGHT):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    set_table_geometry(table, widths)
    style_table(table, font_size=font_size, header_fill=header_fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_paragraph(doc, text: str = "", *, bold_prefix: str | None = None, color: str = INK, italic: bool = False):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_run(paragraph.add_run(bold_prefix), bold=True, color=color)
        set_run(paragraph.add_run(text[len(bold_prefix):]), color=color, italic=italic)
    else:
        set_run(paragraph.add_run(text), color=color, italic=italic)
    return paragraph


def _add_numbering_definition(doc: Document, *, kind: str, marker: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), kind)
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), marker)
    level.append(lvl_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def new_numbering_instance(doc: Document, source_num_id: int) -> int:
    """Create a fresh numbering instance so a new ordered list restarts at 1."""

    numbering = doc.part.numbering_part.element
    source = next(
        item
        for item in numbering.findall(qn("w:num"))
        if int(item.get(qn("w:numId"))) == source_num_id
    )
    abstract_id = source.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(
        (int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))),
        default=0,
    ) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return new_id


def add_list_item(doc, text: str, *, num_id: int, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.2
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))
    paragraph._p.get_or_add_pPr().append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        set_run(paragraph.add_run(bold_prefix), bold=True)
        set_run(paragraph.add_run(text[len(bold_prefix):]))
    else:
        set_run(paragraph.add_run(text))
    return paragraph


def add_callout(doc, label: str, text: str, *, fill: str = INDIGO_TINT, accent: str = INDIGO):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    set_run(paragraph.add_run(label.upper() + "  "), size=9, color=accent, bold=True)
    set_run(paragraph.add_run(text), size=10.5, color=INK)


def add_heading(doc, text: str, level: int = 1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))
    set_run(run, size=8.5, color=MUTED)


def configure_document(doc: Document) -> tuple[int, int]:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after, color in (
        (1, 16, 18, 10, INDIGO),
        (2, 13, 14, 7, INDIGO),
        (3, 12, 10, 5, "1F4D78"),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(header.add_run("Q-TRACE  |  SIH26139 PROJECT GUIDE"), size=8, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(footer.add_run("Research prototype  |  Page "), size=8.5, color=MUTED)
    add_page_number(footer)
    bullet_id = _add_numbering_definition(doc, kind="bullet", marker="•")
    number_id = _add_numbering_definition(doc, kind="decimal", marker="%1.")
    return bullet_id, number_id


def pct(value: float) -> str:
    return f"{100 * float(value):.2f}%"


def load_evidence() -> tuple[dict, dict]:
    return (
        json.loads(CV_PATH.read_text(encoding="utf-8")),
        json.loads(CLASSICAL_PATH.read_text(encoding="utf-8")),
    )


def build() -> Path:
    cv, classical = load_evidence()
    doc = Document()
    bullet_id, number_id = configure_document(doc)
    doc.core_properties.title = "Q-TRACE: SIH26139 Project and Judge Guide"
    doc.core_properties.subject = "Hybrid quantum-classical ML for early disease detection"
    doc.core_properties.author = "Q-TRACE"
    doc.core_properties.keywords = "SIH26139, quantum machine learning, VQC, QSVM, explainability"

    cover_spacer = doc.add_paragraph()
    cover_spacer.paragraph_format.space_after = Pt(92)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(kicker.add_run("SMART INDIA HACKATHON 2026  ·  PROBLEM SIH26139"), size=10, color=TEAL, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    set_run(title.add_run("Q-TRACE"), size=34, color=INDIGO, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    set_run(subtitle.add_run("Hybrid quantum-classical machine learning\nfor early disease signal analysis"), size=16, color="1F4D78")
    strap = doc.add_paragraph()
    strap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    strap.paragraph_format.space_after = Pt(72)
    set_run(strap.add_run("A project explanation, verified evidence record, demo guide, and judge-facing pitch"), size=11, color=MUTED, italic=True)
    add_callout(doc, "Positioning", "A simulator-based research prototype that compares quantum and classical models transparently. It does not claim quantum advantage or clinical readiness.", fill=TEAL_TINT, accent=TEAL)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(26)
    set_run(meta.add_run("Prepared 3 September 2026  |  Hackathon date: 8 September 2026"), size=9.5, color=MUTED, bold=True)
    doc.add_page_break()

    add_heading(doc, "1. The idea in one minute")
    add_paragraph(doc, "Breast-cancer records contain many measurements. Q-TRACE first cleans those measurements, selects four informative features without destroying their clinical names, then asks two different quantum model families and four classical model families to classify the same case. The website always shows quantum and classical results side by side, exposes disagreement, and explains which real measurements influenced each model.")
    add_callout(doc, "Simple analogy", "Think of it as two specialist panels examining the same evidence: one quantum panel with six members, and one classical panel tested under both a powerful 30-feature view and a fair matched 4-feature view.")
    add_heading(doc, "What makes the project credible", 2)
    for item in (
        "Two quantum paradigms: trainable variational circuits and quantum-kernel SVMs.",
        "Four classical algorithms evaluated on all 30 features and the same selected four.",
        "Malignant-positive clinical metrics plus fold-isolated preprocessing.",
        "Real-name SHAP/LIME and saved, configuration-linked evidence.",
    ):
        add_list_item(doc, item, num_id=bullet_id)

    add_heading(doc, "2. What SIH26139 asks for")
    add_paragraph(doc, "The problem asks for a hybrid quantum-classical machine-learning platform for early disease detection. It calls for classical preprocessing and feature engineering, quantum-enhanced classifiers such as QSVM/QNN/VQC, biomedical data ingestion, training, prediction, explainability, classical benchmarking, computational-efficiency evidence, generalization evidence, simulator compatibility, and a path toward near-term hardware.")
    requirement_rows = [
        ["Hybrid architecture", "Train-only classical preprocessing feeds VQC and QSVM quantum paths; classical baselines run beside them.", "Complete"],
        ["Quantum classification", "4 VQCs + angle/4q QSVM + amplitude/2q QSVM; OOB soft-vote ensemble.", "Complete"],
        ["Clinical metrics", "Accuracy, malignant sensitivity/specificity, precision/F1, confusion counts, ROC-AUC.", "Complete"],
        ["Interpretability", "Quantum and classical SHAP + LIME with original feature names.", "Complete"],
        ["Generalization / efficiency", "3-seed reports, 5 isolated folds, measured fit/predict and quantum pipeline time.", "Complete"],
        ["Simulator / hardware path", "Configurable PennyLane device; two simulator paths tested. Real QPU remains future validation.", "Engineering path"],
    ]
    add_table(doc, ["Objective", "Q-TRACE implementation", "Status"], requirement_rows, [2000, 5660, 1700], font_size=8.2, header_fill=INDIGO_TINT)
    add_callout(doc, "Important boundary", "The platform satisfies the software and research-demo scope on WBCD. It does not establish prospective early detection, hospital deployment, regulatory compliance, or real-hardware performance.", fill=AMBER_TINT, accent=AMBER)

    add_heading(doc, "3. Definitions without jargon")
    definitions = [
        ["Biomedical dataset", "A structured collection of measurements and known outcomes. This prototype uses the public Wisconsin Breast Cancer Diagnostic dataset."],
        ["Feature", "One measured property, such as worst radius or mean concave points."],
        ["Label", "The outcome the model learns. In sklearn WBCD, 0 means malignant and 1 means benign."],
        ["Qubit", "The basic information unit of a quantum circuit. A simulator calculates its quantum state on a normal CPU."],
        ["Embedding", "The rule that turns classical feature values into quantum-circuit states."],
        ["VQC / QNN", "A trainable quantum circuit. Classical optimization changes rotation weights to reduce prediction loss."],
        ["QSVM", "A support-vector machine using similarities computed by a quantum feature map. The SVM itself is classical."],
        ["Kernel", "A similarity function. Here it is quantum-state fidelity between two embedded records."],
        ["OOB weighting", "Each ensemble member is scored on rows omitted from its bootstrap sample; lower probability error earns more voting weight."],
        ["Sensitivity", "Of all malignant cases, the fraction correctly flagged malignant. Missing one is a false negative."],
        ["Specificity", "Of all benign cases, the fraction correctly identified benign."],
        ["SHAP / LIME", "Two model-agnostic attribution methods used to cross-check which inputs pushed a prediction toward benign or malignant."],
    ]
    add_table(doc, ["Term", "Meaning in this project"], definitions, [1900, 7460], font_size=8.7, header_fill=TEAL_TINT)

    add_heading(doc, "4. End-to-end flow")
    flow = [
        ["1  INGEST", "2  SPLIT", "3  PREPROCESS", "4  SELECT", "5  MODEL", "6  COMPARE", "7  EXPLAIN"],
        ["30 named values", "Stratified rows", "Impute + scale", "ANOVA top 4", "Quantum + classical", "Equal results", "SHAP + LIME"],
    ]
    add_table(doc, flow[0], [flow[1]], [1337, 1337, 1337, 1337, 1337, 1337, 1338], font_size=7.3, header_fill=INDIGO_TINT)
    steps = (
        "Ingest a held-out record or an exact named 30-feature CSV row.",
        "Split first, preserving untouched test or validation rows.",
        "Fit median imputation and standard scaling only on training rows.",
        "Select four real columns with SelectKBest(f_classif), never PCA.",
        "Scale quantum inputs inside (0, pi) and preserve three label representations.",
        "Run six quantum members plus both classical views; apply frozen OOB weights.",
        "Show both verdicts equally, then compute attribution on request.",
    )
    for step in steps:
        add_list_item(doc, step, num_id=number_id)

    add_heading(doc, "5. Data pipeline and leakage control")
    add_paragraph(doc, "WBCD contains 569 records and 30 numeric features derived from digitized fine-needle-aspirate images of breast masses. Q-TRACE uses the sklearn copy of this benchmark. It is useful for a reproducible hackathon prototype, but it is neither a screening population nor a prospective hospital cohort.")
    add_heading(doc, "Selected clinical signals", 2)
    for name in cv["folds"][0]["selected_feature_names"]:
        add_list_item(doc, name, num_id=bullet_id)
    add_paragraph(doc, "All five cross-validation folds selected the same four names in the retained run. That consistency is informative but not proof they will remain optimal in another cohort.")
    add_heading(doc, "Leakage-safe rule", 2)
    add_callout(doc, "Fold isolation", "For each fold, imputer statistics, standard-scaler parameters, ANOVA scores, selected columns, and quantum min/max scaling are learned only from that fold's training rows. Validation rows are transformed afterward.", fill=TEAL_TINT, accent=TEAL)
    add_heading(doc, "6. The model system")
    add_heading(doc, "Quantum panel: six members", 2)
    quantum_rows = [
        ["VQC", "4 qubits / 3 layers / CNOT", "Primary variational circuit; data re-uploaded each layer"],
        ["VQC", "3 qubits / 2 layers / CNOT", "Shallower top-three-feature diversity"],
        ["VQC", "4 qubits / 2 layers / CZ", "Different entangler"],
        ["VQC", "4 qubits / 3 layers / CZ reversed", "Different encoding order"],
        ["QSVM", "Angle embedding / 4 qubits", "Primary quantum-kernel model"],
        ["QSVM", "Amplitude embedding / 2 qubits", "Four values fit exactly into 2^2 amplitudes; zero-vector safe"],
    ]
    add_table(doc, ["Family", "Configuration", "Role"], quantum_rows, [1450, 3300, 4610], font_size=8.5, header_fill=INDIGO_TINT)
    add_paragraph(doc, "Every VQC uses data re-uploading and sigmoid weight remapping, and trains for at least 100 epochs outside the 20-epoch regression test. The loss converts {-1,+1} labels to {0,1} only inside binary cross-entropy math. PennyLane optimization uses its supported autograd path, avoiding the verified torch/Adam silent-training bug.")
    add_heading(doc, "Classical panel: eight fits", 2)
    add_paragraph(doc, "Logistic Regression, Random Forest, XGBoost, and RBF SVM each run twice: once with all 30 standardized features and once with exactly the four selected features seen by the quantum models. This separates a practical full-information benchmark from a fair input-matched comparison.")
    add_heading(doc, "Voting", 2)
    add_paragraph(doc, "Each quantum member trains on a bootstrap of the quantum training pool. Its probability mean-squared error on out-of-bag rows becomes an inverse-error weight. The six benign probabilities are combined by normalized weighted soft voting; the untouched held-out labels never choose those weights.")

    add_heading(doc, "7. Verified results: classical three-seed report")
    add_paragraph(doc, "All values below are means across seeds 42, 123, and 2026. Malignant is the positive condition. Sensitivity equals malignant recall; the retained JSON also contains per-seed confusion matrices and measured timing.")
    classical_rows = []
    for key, summary in classical["summaries"].items():
        _, view, model = key.split("/")
        metrics = summary["metrics"]
        classical_rows.append([
            ("30 / " if view == "full_feature" else "4 / ") + model.replace("_", " "),
            pct(metrics["accuracy"]["mean"]),
            pct(metrics["precision"]["mean"]),
            pct(metrics["malignant_sensitivity"]["mean"]),
            pct(metrics["specificity"]["mean"]),
            pct(metrics["f1"]["mean"]),
            f"{metrics['roc_auc']['mean']:.3f}",
        ])
    add_table(doc, ["View / model", "Acc", "M-Prec", "Sens", "Spec", "M-F1", "AUC"], classical_rows, [2700, 1110, 1110, 1110, 1110, 1110, 1110], font_size=7.5, header_fill=TEAL_TINT)
    add_callout(doc, "Finding", "Full-feature classical SVM was strongest in this three-seed report at 97.95% mean accuracy. Same-four classical models clustered at 93.27%-93.86%, showing how much of the apparent gap comes from unequal feature access.", fill=TEAL_TINT, accent=TEAL)
    doc.add_page_break()

    add_heading(doc, "8. Verified results: leakage-safe five-fold study")
    add_paragraph(doc, "Each of five stratified folds independently rebuilds preprocessing. Every VQC receives 100 epochs; the quantum pool is deliberately restricted to 20 rows to match the lightweight live configuration. Classical models use the available fold-training rows.")
    cv_rows = []
    for key, summary in cv["summaries"].items():
        paradigm, view, model = key.split("/")
        metrics = summary["metrics"]
        label = ("Q / " if paradigm == "quantum" else "C30 / " if view == "full_feature" else "C4 / ") + model.replace("six_model_oob_ensemble", "six-model ensemble").replace("_", " ")
        cv_rows.append([
            label,
            pct(metrics["accuracy"]["mean"]),
            f"{pct(metrics['accuracy']['range'][0])}-{pct(metrics['accuracy']['range'][1])}",
            pct(metrics["malignant_sensitivity"]["mean"]),
            pct(metrics["specificity"]["mean"]),
            f"{metrics['roc_auc']['mean']:.3f}",
        ])
    add_table(doc, ["Fold model", "Mean acc", "Range", "Sens", "Spec", "AUC"], cv_rows, [3000, 1110, 1800, 1150, 1150, 1150], font_size=7.1, header_fill=LIGHT)
    add_callout(doc, "Honest interpretation", "The six-model ensemble averaged 91.57% accuracy, 83.50% malignant sensitivity, and 96.36% specificity. It did not beat full-feature classical ML. The amplitude QSVM was unstable at 54.61% mean accuracy under the 20-row constraint. Those are limitations to explain, not numbers to hide or tune away after validation.", fill=AMBER_TINT, accent=AMBER)
    doc.add_page_break()
    add_heading(doc, "Computational efficiency", 2)
    timing_rows = []
    for key, summary in classical["summaries"].items():
        _, view, model = key.split("/")
        timing_rows.append([
            ("30 / " if view == "full_feature" else "4 / ") + model.replace("_", " "),
            f"{1000 * summary['mean_fit_seconds']:.2f} ms",
            f"{1000 * summary['mean_predict_seconds']:.2f} ms",
        ])
    ensemble_key = "quantum/same_4_quantum_range/six_model_oob_ensemble"
    timing_rows.append(["Quantum / complete six-model fold pipeline", f"{cv['summaries'][ensemble_key]['mean_fit_seconds']:.2f} s", "included internally"])
    add_table(doc, ["Measured scope", "Mean fit / pipeline", "Mean held-out predict"], timing_rows, [5000, 2180, 2180], font_size=8.0, header_fill=LIGHT)
    add_paragraph(doc, "Timing is hardware- and workload-dependent. The quantum row includes training, OOB weighting, and internal held-out evaluation for the six-member pipeline; it is not directly equivalent to one classical estimator fit.", color=MUTED, italic=True)

    add_heading(doc, "9. Explainability and interface")
    add_heading(doc, "One verdict number, multiple attribution scopes", 2)
    add_paragraph(doc, "The /predict endpoint supplies the only confidence rendered for a patient. The /explain endpoint returns names, SHAP values, LIME values, directions, top features, scope, and timing only. This prevents a fast VQC-only explanation from displaying a second probability that may differ from the six-model ensemble verdict.")
    scope_rows = [
        ["Quantum fast", "Four VQCs", "Default; about 7-30 seconds by device/settings"],
        ["Quantum full", "All six quantum models", "Explicit slow opt-in; 70 seconds or longer"],
        ["Classical full", "30-feature Logistic Regression", "Attribution over all original columns"],
        ["Classical matched", "4-feature Logistic Regression", "Attribution over the identical selected features"],
    ]
    add_table(doc, ["Scope", "Explained model", "UI behavior"], scope_rows, [1900, 3000, 4460], font_size=8.5, header_fill=INDIGO_TINT)
    add_heading(doc, "Judge-facing safeguards", 2)
    for item in (
        "Two equal-size result dials: quantum indigo and classical teal.",
        "Amber disagreement banner when labels differ; neither model is silently chosen.",
        "Circuit rail shows Data -> Select -> Model -> Explain.",
        "CSV schema/range validation warns about out-of-benchmark values without inventing clinical cutoffs.",
        "Long explanation states disclose actual elapsed time and expected duration.",
        "Research-use-only language remains visible throughout.",
    ):
        add_list_item(doc, item, num_id=bullet_id)
    add_heading(doc, "10. Reproducibility and live configuration")
    add_paragraph(doc, "The repository pins Python 3.14 dependencies, includes automated tests and Docker Compose, and exposes model readiness/configuration through /health. The live manifest is artifacts/models/runtime_manifest.json.")
    manifest = json.loads((ROOT / "artifacts" / "models" / "runtime_manifest.json").read_text(encoding="utf-8"))
    manifest_rows = [
        ["Configuration ID", manifest["configuration_id"]],
        ["Dataset", f"{manifest['dataset']} ({manifest['dataset_rows']} rows)"],
        ["Split", f"seed {manifest['split_seed']}, test size {manifest['test_size']}"],
        ["VQC budget", f"{manifest['vqc_epochs']} epochs for every VQC"],
        ["Quantum pool", f"{manifest['quantum_training_limit']} training rows"],
        ["Classical pool", f"{manifest['classical_training_rows']} training rows"],
        ["Quantum device", manifest["quantum_device"]],
        ["Loading mode", manifest["loading_mode"]],
    ]
    add_table(doc, ["Manifest field", "Declared value"], manifest_rows, [2700, 6660], font_size=9, header_fill=TEAL_TINT)
    add_callout(doc, "Do not mix experiments", "The live q20 configuration is identified and reproducible, but it is not the saved q200 Phase 2 benchmark. The website discloses the live settings; the pitch quotes saved benchmark and cross-validation numbers with their own labels.")
    add_heading(doc, "Run locally", 2)
    commands = (
        "py -3.14 -m venv .venv",
        r".\.venv\Scripts\Activate.ps1",
        "python -m pip install -r requirements.txt",
        "docker compose up --build   (or run uvicorn backend.main:app)",
        "Open http://127.0.0.1:8000/ and wait for Model service ready.",
    )
    number_id = new_numbering_instance(doc, number_id)
    for command in commands:
        add_list_item(doc, command, num_id=number_id)

    add_heading(doc, "11. How to approach the judges")
    add_heading(doc, "Opening: 30 seconds", 2)
    add_callout(doc, "Suggested words", "Q-TRACE is an explainable hybrid quantum-classical research platform for SIH26139. We compare two quantum paradigms - four trainable VQCs and two quantum-kernel SVMs - against four classical algorithms on both full and matched features. Every patient shows both results, every metric treats malignant disease as positive, and every claim is tied to leakage-safe evidence.", fill=INDIGO_TINT, accent=INDIGO)
    add_heading(doc, "Five-minute demonstration", 2)
    demo_steps = (
        "Show the circuit rail and explain that preprocessing is classical while classification is compared across paradigms.",
        "Select a held-out record, point to the four real selected measurements, and run hybrid analysis.",
        "Read both equal-size dials. If they disagree, use the amber banner as evidence of transparency.",
        "Open per-model quantum details to show four VQCs, angle QSVM, amplitude QSVM, and OOB weights.",
        "Generate fast quantum attribution, then switch to a classical attribution scope. Emphasize real feature names and one verdict number.",
        "Scroll to five-fold sensitivity/specificity evidence and close with limitations plus the hardware-validation path.",
    )
    number_id = new_numbering_instance(doc, number_id)
    for item in demo_steps:
        add_list_item(doc, item, num_id=number_id)
    add_heading(doc, "Questions to welcome", 2)
    qa = [
        ["Why quantum if classical is more accurate?", "The project investigates two quantum learning mechanisms under a transparent baseline. Current data favors full-feature classical ML; the value is the reproducible comparison, ensemble stability study, and hardware-ready circuit abstraction - not a fabricated advantage claim."],
        ["Why only four quantum features?", "Near-term circuits are constrained by qubits and kernel cost. ANOVA selection preserves interpretable clinical names and enables a fair four-feature classical comparison."],
        ["Why is amplitude QSVM weak?", "Four values fit naturally into two qubits, but normalization discards magnitude and the 20-row fold pool is severe. Its instability is measured and down-weighted rather than hidden."],
        ["Is this early detection?", "It is an early-disease ML platform prototype evaluated on diagnostic breast-mass data. Prospective screening or earlier-stage benefit would require an external longitudinal cohort."],
        ["Does it run on real quantum hardware?", "The circuit device is configurable and simulator paths are tested. A real QPU still needs a provider plugin, shot/noise settings, gate validation, credentials, and a separate measured report."],
    ]
    add_table(doc, ["Judge question", "Defensible answer"], qa, [3100, 6260], font_size=8.1, header_fill=AMBER_TINT)
    add_heading(doc, "12. What to improve next")
    add_paragraph(doc, "The most impressive next steps are validation steps, not extra animations or unsupported model claims.")
    next_steps = (
        "External-cohort validation: test on a separately sourced breast-cancer cohort and lock preprocessing/model settings before looking at outcomes.",
        "Clinical thresholding: select a high-sensitivity operating point on training data, then report its specificity and calibration on untouched data.",
        "Subgroup analysis: assess false negatives and calibration across clinically relevant groups when lawful, appropriate metadata is available.",
        "Noise-aware simulation: run the same circuits under documented noise channels and report accuracy/latency changes.",
        "Small real-QPU appendix: execute a narrowly scoped circuit on supported hardware and compare simulator/noisy/hardware outputs without claiming advantage.",
        "Multimodal research: add imaging or genomic representations only after the core tabular pipeline remains leakage-safe and interpretable.",
    )
    number_id = new_numbering_instance(doc, number_id)
    for item in next_steps:
        add_list_item(doc, item, num_id=number_id)
    add_heading(doc, "Known limitations", 2)
    for item in (
        "One public benchmark dataset; no prospective or external clinical validation.",
        "The live and five-fold quantum pool is 20 rows, while the separate headline Phase 2 benchmark uses 200.",
        "Quantum execution is simulated on CPU; no quantum speedup or hardware robustness is demonstrated.",
        "The OOB weights are model-level estimates inside each fold; preprocessing is shared within the fold's training partition.",
        "SHAP and LIME explain model behavior, not biological causality or clinical action.",
        "The local demo lacks authentication, HTTPS termination, hospital integration, and a protected-health-data workflow.",
    ):
        add_list_item(doc, item, num_id=bullet_id)

    doc.add_page_break()
    add_heading(doc, "13. Verification checklist")
    checks = (
        "Run the full pytest suite and the VQC cost-decrease regression.",
        "Confirm pip check, Python compilation, and frontend JavaScript syntax.",
        "Start Docker from the current commit and wait for six quantum plus eight classical fits.",
        "Test benchmark selection, exact named CSV intake, agreement, disagreement, fast quantum explanation, and both classical explanation scopes.",
        "Confirm the evidence table loads the retained five-fold metrics.",
        "Rehearse on the presentation machine and keep a labeled backup recording.",
        "Never upload credentials or identifiable patient records to GitHub.",
    )
    for item in checks:
        add_list_item(doc, item, num_id=bullet_id)

    add_heading(doc, "14. Evidence and sources")
    sources = (
        "SIH26139 problem listing mirror (accessed 3 September 2026): https://sih2026.vuce.in/ps/SIH26139",
        "UCI Breast Cancer Wisconsin (Diagnostic): https://archive.ics.uci.edu/dataset/17/breast+cancer+wisconsin+diagnostic",
        "scikit-learn breast cancer dataset documentation: https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_breast_cancer.html",
        "PennyLane QNode documentation: https://docs.pennylane.ai/en/stable/code/api/pennylane.QNode.html",
        "PennyLane kernels demonstration: https://pennylane.ai/qml/demos/tutorial_kernels_module",
        "Repository evidence: artifacts/models/phase2_benchmark_100epochs_200pool.md",
        "Repository evidence: artifacts/evaluation/classical_three_seed_metrics.json",
        "Repository evidence: artifacts/evaluation/five_fold_cross_validation.json",
        "Project contracts: prd.md, architecture.md, decisions.md, roadmap.md, and DEMO_GUIDE.md",
    )
    number_id = new_numbering_instance(doc, number_id)
    for source in sources:
        add_list_item(doc, source, num_id=number_id)
    add_paragraph(doc, "This guide distinguishes implemented evidence from future clinical and hardware validation. Numbers are rounded only for readability; machine-readable artifacts retain full precision.", color=MUTED, italic=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
